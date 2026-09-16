"""PowerPoint extraction tests."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.ingestion.pptx_parser import PptxParser


def _make_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Network Security"
    slide.placeholders[1].text = "Firewalls filter traffic between trust zones."
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Mention DMZ and VLAN segmentation in the lab."
    prs.save(path)


def test_pptx_extraction_preserves_slides_and_notes(tmp_path: Path) -> None:
    pptx = tmp_path / "lecture.pptx"
    _make_pptx(pptx)
    extracted = PptxParser().parse(pptx, "Year2/Networking/lecture.pptx")
    assert extracted.file_type == "pptx"
    assert extracted.metadata.get("slide_count") == 1
    pages = {b.page for b in extracted.blocks}
    assert 1 in pages
    blob = "\n".join(b.text for b in extracted.blocks)
    assert "Network Security" in blob
    assert "Firewalls filter traffic" in blob
    assert "VLAN segmentation" in blob
    assert any(b.block_type == "notes" for b in extracted.blocks)
