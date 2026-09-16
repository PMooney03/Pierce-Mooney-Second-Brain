"""PowerPoint extraction with python-pptx + OCR on embedded slide images."""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.picture import Picture
from pptx.table import Table

from app.config import get_settings
from app.database.models import ExtractedBlock, ExtractedDocument
from app.ingestion.base_parser import DocumentParser
from app.ingestion.ocr_engine import OcrError, ocr_image
from app.logging_config import get_logger

logger = get_logger(__name__)


class PptxParseError(Exception):
    """Raised for unreadable / empty PowerPoint files."""


def _table_to_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        cells = [c for c in cells if c]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _shape_text(shape: BaseShape) -> str:
    if shape.has_text_frame:
        parts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
        return "\n".join(parts)
    if shape.has_table:
        return _table_to_text(shape.table)
    return ""


def _iter_shapes(shapes) -> list[BaseShape]:
    out: list[BaseShape] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(_iter_shapes(shape.shapes))
        else:
            out.append(shape)
    return out


def _slide_title(slide) -> str | None:
    try:
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()
            return title or None
    except Exception:  # noqa: BLE001
        pass
    return None


def _ocr_picture(shape: Picture) -> str:
    settings = get_settings()
    if not settings.ocr_enabled:
        return ""

    try:
        blob = shape.image.blob
    except Exception as exc:  # noqa: BLE001
        logger.debug("Could not read picture blob: %s", exc)
        return ""

    max_bytes = int(settings.ocr_max_file_mb * 1024 * 1024)
    if len(blob) > max_bytes:
        logger.debug("Skipping large slide image (%s bytes)", len(blob))
        return ""

    try:
        from PIL import Image
        import numpy as np

        with Image.open(io.BytesIO(blob)) as img:
            width, height = img.size
            if min(width, height) < settings.ocr_min_image_side:
                return ""
            rgb = img.convert("RGB")
            text = ocr_image(np.asarray(rgb))
    except OcrError as exc:
        logger.debug("Slide image OCR failed: %s", exc)
        return ""
    except Exception as exc:  # noqa: BLE001
        logger.debug("Slide image open failed: %s", exc)
        return ""

    return (text or "").strip()


class PptxParser(DocumentParser):
    extensions = {".pptx"}

    def parse(self, path: Path, relative_path: str) -> ExtractedDocument:
        try:
            presentation = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise PptxParseError(f"Cannot open PPTX: {exc}") from exc

        blocks: list[ExtractedBlock] = []
        ocr_images = 0

        for slide_index, slide in enumerate(presentation.slides, start=1):
            heading = _slide_title(slide)
            slide_texts: list[str] = []

            for shape in _iter_shapes(slide.shapes):
                if isinstance(shape, Picture):
                    ocr_text = _ocr_picture(shape)
                    if ocr_text:
                        ocr_images += 1
                        blocks.append(
                            ExtractedBlock(
                                text=ocr_text,
                                page=slide_index,
                                heading=heading,
                                block_type="ocr",
                            )
                        )
                    continue

                text = _shape_text(shape)
                if not text:
                    continue
                # Avoid duplicating the title if it appears in a title placeholder.
                if heading and text.strip() == heading and len(slide_texts) == 0:
                    blocks.append(
                        ExtractedBlock(
                            text=text,
                            page=slide_index,
                            heading=heading,
                            block_type="heading",
                        )
                    )
                    continue
                slide_texts.append(text)
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        page=slide_index,
                        heading=heading,
                        block_type="paragraph",
                    )
                )

            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes_text:
                blocks.append(
                    ExtractedBlock(
                        text=notes_text,
                        page=slide_index,
                        heading=heading,
                        block_type="notes",
                    )
                )

        if not blocks:
            raise PptxParseError("Empty PPTX (no extractable text or OCR content)")

        return ExtractedDocument(
            filepath=relative_path.replace("\\", "/"),
            filename=path.name,
            file_type="pptx",
            blocks=blocks,
            metadata={
                "slide_count": len(presentation.slides),
                "ocr_images": ocr_images,
            },
        )
